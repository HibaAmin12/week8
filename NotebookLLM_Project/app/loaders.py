"""File loaders. Each loader returns a list of "units": the natural pieces of a file.

A unit is a dict: {label, text, heading, mergeable}
  label     what the citation shows ("Page 3", "Slide 2", "Section: Methods", "Rows 1-25")
  text      the unit's full text (the viewer highlights chunks inside it)
  heading   optional heading, prepended to chunks when embedding
  mergeable tiny units (headings with little text) may merge into the next unit
"""
import csv
import json
import os
import re

from . import config


def _clean(text: str) -> str:
    text = text.replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)  # single newlines are just line wraps
    return text.strip()


def _unit(label, text, heading="", mergeable=False):
    return {"label": label, "text": text, "heading": heading, "mergeable": mergeable}


# ---------------- OCR ----------------
def _ocr_image(img) -> str:
    if not config.OCR_ENABLED:
        return ""
    try:
        import pytesseract
        return pytesseract.image_to_string(img, lang=config.OCR_LANG)
    except Exception as exc:  # tesseract missing, language pack missing, etc.
        print(f"[ocr] failed: {exc}")
        return ""


def _load_pdf(path):
    import fitz  # PyMuPDF
    units = []
    with fitz.open(path) as pdf:
        for i, page in enumerate(pdf, 1):
            text = _clean(page.get_text())
            if len(text) < config.OCR_MIN_CHARS and config.OCR_ENABLED:
                from PIL import Image
                pix = page.get_pixmap(dpi=config.OCR_DPI)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                text = _clean(_ocr_image(img))
            if text:
                units.append(_unit(f"Page {i}", text))
    return units


def _load_image(path):
    from PIL import Image
    text = _clean(_ocr_image(Image.open(path)))
    return [_unit("Image", text)] if text else []


# ---------------- Word ----------------
def _load_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    units, heading, buf = [], "Introduction", []

    def flush():
        text = _clean("\n\n".join(buf))
        if text:
            units.append(_unit(f"Section: {heading}", text, heading, mergeable=True))

    for child in d.element.body.iterchildren():
        if child.tag.endswith("}p"):
            p = Paragraph(child, d)
            txt = p.text.strip()
            if not txt:
                continue
            if (p.style.name or "").lower().startswith(("heading", "title")):
                flush()
                buf, heading = [], txt
            else:
                buf.append(txt)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, d).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    buf.append(" | ".join(cells))
    flush()
    return units


# ---------------- PowerPoint ----------------
def _load_pptx(path):
    from pptx import Presentation
    units = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        text = _clean("\n\n".join(p for p in parts if p.strip()))
        if text:
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
            units.append(_unit(f"Slide {i}", text, title))
    return units


# ---------------- Spreadsheets ----------------
def _rows_to_units(sheet_name, rows):
    """Group rows; every row is written as 'column: value' so column names appear in every chunk."""
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    header = [h or f"col{i + 1}" for i, h in enumerate(header)]
    units, body = [], rows[1:]
    for start in range(0, len(body), config.ROWS_PER_UNIT):
        group = body[start:start + config.ROWS_PER_UNIT]
        lines = ["; ".join(f"{header[i]}: {str(v).strip()}"
                           for i, v in enumerate(r[:len(header)]) if v is not None and str(v).strip())
                 for r in group]
        label = f"{sheet_name} rows {start + 1}-{start + len(group)}"
        units.append(_unit(label, "\n\n".join(lines), sheet_name))
    return units


def _load_csv(path):
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        return _rows_to_units("CSV", list(csv.reader(f)))


def _load_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    units = []
    for ws in wb.worksheets:
        units += _rows_to_units(ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
    return units


# ---------------- Text-like ----------------
PAGE_MARK = re.compile(r"\[\[PAGE\s+(\d+)\]\]", re.I)


def _read_text(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def _split_blocks(text, size):
    """Split long text into units of about `size` characters on paragraph boundaries."""
    parts, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        if cur and len(cur) + len(para) > size:
            parts.append(cur)
            cur = ""
        cur = (cur + "\n\n" + para).strip()
    if cur:
        parts.append(cur)
    return parts


def _load_txt(path):
    raw = _read_text(path)
    marks = list(PAGE_MARK.finditer(raw))
    if len(marks) >= 2:  # text exported with [[PAGE n]] markers
        units = []
        for i, m in enumerate(marks):
            end = marks[i + 1].start() if i + 1 < len(marks) else len(raw)
            text = _clean(raw[m.end():end])
            if text:
                units.append(_unit(f"Page {m.group(1)}", text))
        return units
    return [_unit(f"Part {i}", t) for i, t in enumerate(_split_blocks(_clean(raw), config.TEXT_UNIT_CHARS), 1)]


def _load_md(path):
    units, heading, buf, in_code = [], "Introduction", [], False

    def flush():
        text = _clean("\n\n".join(buf))
        if text:
            units.append(_unit(f"Section: {heading}", text, heading, mergeable=True))

    for line in _read_text(path).splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
        m = None if in_code else re.match(r"^#{1,6}\s+(.*)", line)
        if m:
            flush()
            buf, heading = [], m.group(1).strip()
        else:
            buf.append(line)
    flush()
    return units


def _load_json(path):
    text = json.dumps(json.loads(_read_text(path)), indent=1, ensure_ascii=False)
    return [_unit(f"Part {i}", t) for i, t in enumerate(_split_blocks(text, config.TEXT_UNIT_CHARS), 1)]


LOADERS = {".pdf": _load_pdf, ".docx": _load_docx, ".pptx": _load_pptx, ".xlsx": _load_xlsx,
           ".csv": _load_csv, ".md": _load_md, ".txt": _load_txt, ".json": _load_json,
           ".png": _load_image, ".jpg": _load_image, ".jpeg": _load_image}


def load(path: str) -> list:
    ext = os.path.splitext(path)[1].lower()
    if ext not in LOADERS:
        raise ValueError(f"Unsupported file type: {ext}")
    return LOADERS[ext](path)
